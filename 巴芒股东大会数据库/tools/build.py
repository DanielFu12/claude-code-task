#!/usr/bin/env python3
"""由 entries/*.txt + 章节逐字稿生成全部交付物。

用法：
    python3 tools/build.py CHAPTERS.jsonl [OUT_DIR]

条目文件格式（entries/YYYY.txt）：每条以 “@ 章节ID” 开头，字段一行一个——
  t 标题 / s 发言人(B 巴菲特 M 芒格 A 阿贝尔 J 贾因，可组合) / th 主题(分号分隔) / k 概念簇 /
  c 核心观点 / q 原话锚点“起始短语 ||| 结束短语”(可多行) / zh 中文翻译 / l 投资启示 /
  r 评分 深度/实用/启发(各1-5) / rel 相关章节(可选)。详见 tools/entries.py。

时间戳说明：逐字稿没有时间码。本脚本按“该句在本场逐字稿中的累计词数 ÷ 语速”估算秒数。
语速默认 2.1 词/秒（由 2020、2024 已知时长场次反推），已知时长的场次用实际时长折算。
因此所有秒数都是估算值，深链在估算区间前后各留 60 秒缓冲。
"""
import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from entries import CLUSTERS, THEMES, load_chapters, parse_entries, validate  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RATE = 2.1          # 词/秒，未知时长场次的默认语速
BUFFER = 60         # 深链前后缓冲秒数
# 已知时长（秒）的场次：来自 CNBC 页面标注
KNOWN_DURATION = {
    (2020, 'F'): 2 * 3600 + 6 * 60,
    (2024, 'M'): 2 * 3600 + 55 * 60,
    (2024, 'A'): 1 * 3600 + 57 * 60,
}
SPEAKERS = {'B': '巴菲特', 'M': '芒格', 'A': '阿贝尔', 'J': '贾因'}
SESSION_ZH = {'M': '上午场', 'A': '下午场', 'F': '全场'}
# 股东大会日期。带 * 的来自 CNBC 档案页标注，其余依据公开资料，未在档案页逐一核实。
MEETING_DATE = {
    1994: '1994-04-25*', 1995: '1995-05-01', 1996: '1996-05-06*', 1997: '1997-05-05*',
    1998: '1998-05-04*', 1999: '1999-05-03', 2000: '2000-04-29', 2001: '2001-04-28',
    2002: '2002-05-04*', 2003: '2003-05-03', 2004: '2004-05-01', 2005: '2005-04-30',
    2006: '2006-05-06', 2007: '2007-05-05*', 2008: '2008-05-03*', 2009: '2009-05-02',
    2010: '2010-05-01', 2011: '2011-04-30*', 2012: '2012-05-05*', 2013: '2013-05-04*',
    2014: '2014-05-03', 2015: '2015-05-02', 2016: '2016-04-30', 2017: '2017-05-06*',
    2018: '2018-05-05', 2019: '2019-05-04', 2020: '2020-05-02*', 2021: '2021-05-01*',
    2022: '2022-04-30*', 2023: '2023-05-06*', 2024: '2024-05-04*', 2025: '2025-05-03*',
    2026: '2026-05-02*',
}


def hms(sec):
    sec = max(0, int(round(sec)))
    return f'{sec // 3600}:{sec % 3600 // 60:02d}:{sec % 60:02d}'


def load_urls(path):
    rows = defaultdict(list)
    lines = path.read_text(encoding='utf-8').strip().split('\n')
    for ln in lines[1:]:
        parts = ln.split('\t') + [''] * 5
        year, kind, title, url, note = parts[:5]
        rows[(int(year), kind)].append({'title': title, 'url': url, 'note': note})
    return rows


def session_video(urls, year, sess):
    """章节场次 → CNBC 视频页。2020/2021 用 Part 1/Part 2。"""
    if year == 2020:
        kind = 'part1'
    elif year == 2021:
        kind = 'part1' if sess == 'M' else 'part2'
    else:
        kind = {'M': 'morning', 'A': 'afternoon'}[sess]
    hit = urls.get((year, kind))
    return (hit[0]['url'], hit[0]['title']) if hit else (None, None)


def deep_link(url, start, end):
    if not url:
        return None
    return f'{url}?amp=&start={int(start)}&amp=&end={int(math.ceil(end))}'


def enrich(entries, chapters, urls):
    session_words = {}
    for ch in chapters.values():
        session_words[(ch['year'], ch['id'][5])] = ch['session_words']
    seen = Counter()
    out = []
    for e in entries:
        ch = chapters[e['ch']]
        year, sess = ch['year'], ch['id'][5]
        seen[e['ch']] += 1
        eid = e['ch'] if seen[e['ch']] == 1 else f"{e['ch']}-{seen[e['ch']]}"
        dur = KNOWN_DURATION.get((year, sess))
        rate = session_words[(year, sess)] / dur if dur else RATE
        quotes = []
        for q in e['quotes']:
            s = (ch['word_start'] + q['word_offset']) / rate
            quotes.append({'text': q['text'], 'start': s, 'end': s + q['words'] / rate,
                           'words': q['words']})
        start = min(q['start'] for q in quotes)
        end = max(q['end'] for q in quotes)
        url, vtitle = session_video(urls, year, sess)
        d, p, i = (int(x) for x in e['r'].split('/'))
        out.append({
            'id': eid, 'chapter': e['ch'], 'year': year, 'session': SESSION_ZH[sess],
            'chapter_no': ch['chapter'], 'chapter_title': ch['title'],
            'title': e['t'], 'speakers': '、'.join(SPEAKERS[c] for c in e['s'] if c in SPEAKERS),
            'themes': [t.strip() for t in e['th'].split(';') if t.strip()],
            'cluster': e['k'], 'cluster_name': CLUSTERS[e['k']],
            'core': e['c'], 'quote': ' […] '.join(q['text'] for q in quotes),
            'quotes': quotes, 'zh': e['zh'], 'lesson': e['l'],
            'depth': d, 'practical': p, 'insight': i, 'score': d + p + i,
            'start': start, 'end': end,
            'start_hms': hms(start), 'end_hms': hms(end),
            'video': url, 'video_title': vtitle,
            'link': deep_link(url, max(0, start - BUFFER), end + BUFFER),
            'rel_chapters': e['rel_ids'], 'file': e['file'],
        })
    # 相关条目：同概念簇（按年份远近取 4 条）+ 人工标注的相关章节
    by_cluster = defaultdict(list)
    by_chapter = defaultdict(list)
    for x in out:
        by_cluster[x['cluster']].append(x)
        by_chapter[x['chapter']].append(x)
    for x in out:
        peers = sorted((y for y in by_cluster[x['cluster']] if y is not x),
                       key=lambda y: (abs(y['year'] - x['year']), -y['score']))
        rel = [y['id'] for y in peers[:4]]
        for c in x['rel_chapters']:
            for y in by_chapter.get(c, []):
                if y['id'] not in rel and y is not x:
                    rel.append(y['id'])
        x['related'] = rel
    return out


def hash_id(s):
    """同分条目的稳定排序：不按年份先后，避免偏向早年。"""
    import hashlib
    return hashlib.md5(s.encode()).hexdigest()


def pick_top(items, n=100, per_year=4, per_cluster=5):
    """评分优先，同时限制单一年份/概念簇的条数，保证跨年份、跨主题的多样性。"""
    ranked = sorted(items, key=lambda x: (-x['score'], -x['insight'], -x['depth'],
                                          -x["practical"], hash_id(x["id"])))
    chosen, yc, cc = [], Counter(), Counter()
    for limit_y, limit_c in ((per_year, per_cluster), (per_year + 2, per_cluster + 2), (99, 99)):
        for x in ranked:
            if len(chosen) >= n:
                break
            if x in chosen or yc[x['year']] >= limit_y or cc[x['cluster']] >= limit_c:
                continue
            chosen.append(x)
            yc[x['year']] += 1
            cc[x['cluster']] += 1
    chosen.sort(key=lambda x: (-x['score'], -x['insight'], -x['depth'], x['year']))
    for rank, x in enumerate(chosen, 1):
        x['top_rank'] = rank
    return chosen


def catalog_rows(urls, chapters):
    stats = defaultdict(lambda: {'chapters': 0, 'words': 0, 'sessions': set()})
    for ch in chapters.values():
        s = stats[ch['year']]
        s['chapters'] += 1
        s['words'] += sum(len(p['text'].split()) for p in ch['paragraphs'])
        s['sessions'].add(ch['id'][5])
    rows = []
    for year in range(1994, 2027):
        meet = urls.get((year, 'meeting'), [{}])[0]
        st = stats.get(year)
        vids = []
        for kind, label in (('morning', '上午场'), ('afternoon', '下午场'), ('part1', 'Part 1'),
                            ('part2', 'Part 2'), ('highlights', '精华剪辑'), ('highlights_zh', '精华剪辑(普通话)')):
            for v in urls.get((year, kind), []):
                vids.append({'label': label, 'title': v['title'], 'url': v['url'], 'note': v['note']})
        clips = [{'label': '单段视频', 'title': v['title'], 'url': v['url'], 'note': v['note']}
                 for v in urls.get((year, 'clip'), [])]
        date = MEETING_DATE.get(year, '')
        notes_all = ' '.join(v['note'] for k, vs in urls.items() if k[0] == year for v in vs)
        if date and not date.endswith('*') and date in notes_all:
            date += '*'
        if st:
            sess = '+'.join(SESSION_ZH[s] for s in sorted(st['sessions'], key='MAF'.index))
            cover = f"{st['chapters']} 章 / 约 {st['words']:,} 词（{sess}）"
        else:
            cover = '无逐字稿'
        gap = {
            2019: '逐字稿缺下午场', 2020: '逐字稿仅含 Part 1（演讲与正式会议），缺 Part 2 问答',
            2025: '无逐字稿；仅收录视频链接', 2026: '无逐字稿；会议场次视频链接未核实，仅有精华剪辑与单段视频',
        }.get(year, '')
        if year == 2025:
            gap += '；下午场视频链接未核实'
        rows.append({'year': year, 'date': date.rstrip('*'), 'date_verified': date.endswith('*'),
                     'page': meet.get('url', ''), 'note': meet.get('note', ''),
                     'videos': vids, 'clips': clips, 'transcript': cover, 'gap': gap})
    return rows


def write_markdown(out, items, top, catalog, notes):
    out.mkdir(parents=True, exist_ok=True)
    # 1) 视频目录
    L = ['# 伯克希尔股东大会视频目录（1994–2026）', '',
         '来源：CNBC Warren Buffett Archive（buffett.cnbc.com）。CNBC 档案收录 1994 年起的全部股东大会完整视频。',
         '“逐字稿覆盖”指本库用于逐字核对原话的 CNBC 逐字稿（GitHub 镜像 pzponge/Yestoday）。',
         '日期后无“✓”者未在档案页逐一核实。', '']
    for r in catalog:
        L.append(f"## {r['year']} 年" + (f"（{r['date']}{' ✓' if r['date_verified'] else ''}）" if r['date'] else ''))
        if r['page']:
            L.append(f"- 会议页：<{r['page']}>" + (f"　{r['note']}" if r['note'] else ''))
        for v in r['videos']:
            L.append(f"- {v['label']}：[{v['title']}]({v['url']})" + (f"　{v['note']}" if v['note'] else ''))
        for v in r['clips']:
            L.append(f"- 单段：[{v['title']}]({v['url']})")
        L.append(f"- 逐字稿覆盖：{r['transcript']}" + (f"；**{r['gap']}**" if r['gap'] else ''))
        n = sum(1 for x in items if x['year'] == r['year'])
        L.append(f"- 本库收录精华：{n} 条")
        L.append('')
    (out / '视频目录.md').write_text('\n'.join(L), encoding='utf-8')

    # 2) 按主题的知识库 + 思想演变
    L = ['# 巴菲特与芒格股东大会精华知识库（按主题）', '',
         f'共 {len(items)} 条精华，覆盖 1994–2024 年。每条给出 CNBC 章节、估算时间与视频深链。',
         '时间为按逐字稿词数位置估算，误差通常在数分钟内；深链已在前后各留 60 秒缓冲。', '',
         '## 目录', '']
    for th in THEMES:
        L.append(f"- [{th}](#{th})（{sum(1 for x in items if th in x['themes'])} 条）")
    L += ['- [思想演变（按概念簇）](#思想演变按概念簇)', '']
    for th in THEMES:
        L += [f'## {th}', '']
        group = sorted((x for x in items if th in x['themes']), key=lambda x: (x['year'], x['id']))
        for x in group:
            L.append(f"### {x['year']} · {x['title']}")
            L.append(f"- 发言人：{x['speakers']}　｜　概念簇：{x['cluster_name']}　｜　评分 {x['depth']}/{x['practical']}/{x['insight']}")
            L.append(f"- 核心观点：{x['core']}")
            L.append(f"- 投资启示：{x['lesson']}")
            link = f"[{x['session']} {x['start_hms']}–{x['end_hms']}（估算）]({x['link']})" if x['link'] else f"{x['session']} {x['start_hms']}–{x['end_hms']}（估算）"
            L.append(f"- 出处：CNBC {x['chapter']}「{x['chapter_title']}」，{link}")
            L.append('')
    L += ['## 思想演变（按概念簇）', '',
          '同一思想在不同年份反复出现。下面按概念簇把相关条目按时间排开，便于看到观点如何延续、补充或修正。', '']
    for key, name in CLUSTERS.items():
        group = sorted((x for x in items if x['cluster'] == key), key=lambda x: (x['year'], x['id']))
        if not group:
            continue
        L += [f'### {name}（{len(group)} 条，{group[0]["year"]}–{group[-1]["year"]}）', '']
        if notes.get(key):
            L += [notes[key], '']
        for x in group:
            L.append(f"- **{x['year']}** {x['title']}（{x['speakers']}；{x['chapter']}）")
        L.append('')
    (out / '知识库_按主题.md').write_text('\n'.join(L), encoding='utf-8')

    # 3) Top 100
    L = ['# Top 100 精华片段', '',
         '评选：按“思想深度 + 投资实用性 + 启发性”（各 1–5 分）总分排序，同分时启发性、深度优先；',
         '为保证覆盖面，单一年份最多 4 条、单一概念簇最多 5 条（不足 100 条时逐步放宽）。',
         '时间为估算值；深链在估算区间前后各留 60 秒缓冲。', '']
    for x in top:
        L.append(f"## {x['top_rank']}. {x['title']}")
        L.append(f"**{x['year']} · {x['speakers']} · {x['cluster_name']} · 评分 {x['score']}（{x['depth']}/{x['practical']}/{x['insight']}）**")
        L.append('')
        L.append(f"- 核心观点：{x['core']}")
        L.append(f"- 原话：> {x['quote']}")
        L.append(f"- 中文翻译：{x['zh']}")
        L.append(f"- 投资启示：{x['lesson']}")
        L.append(f"- 视频：{x['video_title'] or '（无场次视频链接）'}")
        L.append(f"- 时间（估算）：{x['session']} {x['start_hms']}–{x['end_hms']}" + (f"　[跳转]({x['link']})" if x['link'] else ''))
        L.append(f"- CNBC 章节：{x['chapter']}「{x['chapter_title']}」")
        L.append('')
    (out / 'Top100.md').write_text('\n'.join(L), encoding='utf-8')


def write_excel(path, items, top, catalog, notes):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill('solid', fgColor='1F3A5F')
    head_font = Font(bold=True, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='top')

    def sheet(ws, headers, rows, widths, link_cols=()):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(vertical='center', wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
            for idx in link_cols:
                c = row[idx]
                if c.value and str(c.value).startswith('http'):
                    c.hyperlink = c.value
                    c.font = Font(color='1155CC', underline='single')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = '说明'
    info = [
        ['巴菲特与芒格股东大会精华数据库'],
        [f'收录 {len(items)} 条精华（1994–2024），Top 100 另见工作表。'],
        ['原话：由程序按起止短语从 CNBC 逐字稿中原样截取，只删去“(Laughter)”等现场标注；换发言人处插入 [Munger:] 等标记。'],
        ['中文翻译、核心观点、投资启示、评分：本库编写，非官方。'],
        ['时间戳：逐字稿无时间码，按“本场累计词数 ÷ 语速（默认 2.1 词/秒）”估算，误差通常数分钟；深链前后各留 60 秒缓冲。'],
        ['视频链接：CNBC Warren Buffett Archive 场次页面；2020 年对应 Part 1，2021 年上午段/下午段对应 Part 1/Part 2。'],
        ['缺口：2019 下午场、2020 问答部分、2025 与 2026 年无逐字稿，未收录精华；2025 下午场与 2026 场次视频链接未核实。'],
        ['评分：思想深度 / 投资实用性 / 启发性，各 1–5 分，主观评定。'],
        ['筛选：在“精华数据库”表头使用筛选按钮，可按年份、发言人、主题、概念簇、评分筛选。'],
    ]
    for r in info:
        ws.append(r)
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 120
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = wrap

    rows = []
    for r in catalog:
        sess = {v['label']: v['url'] for v in r['videos']}
        rows.append([r['year'], r['date'] + ('' if r['date_verified'] or not r['date'] else '（未核实）'),
                     r['page'], sess.get('上午场') or sess.get('Part 1', ''),
                     sess.get('下午场') or sess.get('Part 2', ''), sess.get('精华剪辑', ''),
                     len(r['clips']), r['transcript'], r['gap'], r['note'],
                     sum(1 for x in items if x['year'] == r['year'])])
    sheet(wb.create_sheet('视频目录'),
          ['年份', '会议日期', 'CNBC 会议页', '上午场 / Part 1', '下午场 / Part 2', '精华剪辑',
           '单段视频数', '逐字稿覆盖', '缺口', '备注', '本库条数'],
          rows, [7, 14, 40, 50, 50, 50, 8, 26, 30, 30, 8], link_cols=(2, 3, 4, 5))

    cols = ['编号', '年份', '场次', '发言人', '主题', '概念簇', '标题', '核心观点', '原话（英文）',
            '中文翻译', '投资启示', '深度', '实用', '启发', '总分', '开始（估算）', '结束（估算）',
            '视频链接', '跳转链接（估算）', 'CNBC 章节', '章节标题', '相关条目', 'Top100 名次']
    rows = [[x['id'], x['year'], x['session'], x['speakers'], '；'.join(x['themes']), x['cluster_name'],
             x['title'], x['core'], x['quote'], x['zh'], x['lesson'], x['depth'], x['practical'],
             x['insight'], x['score'], x['start_hms'], x['end_hms'], x['video'], x['link'],
             x['chapter'], x['chapter_title'], '，'.join(x['related']), x.get('top_rank')]
            for x in items]
    widths = [11, 7, 8, 10, 18, 18, 30, 60, 70, 70, 40, 6, 6, 6, 6, 10, 10, 30, 30, 10, 30, 30, 8]
    sheet(wb.create_sheet('精华数据库'), cols, rows, widths, link_cols=(17, 18))

    rows = [[x['top_rank'], x['id'], x['year'], x['speakers'], x['cluster_name'], x['title'], x['core'],
             x['quote'], x['zh'], x['lesson'], x['score'], f"{x['depth']}/{x['practical']}/{x['insight']}",
             f"{x['start_hms']}–{x['end_hms']}", x['link']] for x in top]
    sheet(wb.create_sheet('Top100'),
          ['名次', '编号', '年份', '发言人', '概念簇', '标题', '核心观点', '原话', '中文翻译', '投资启示',
           '总分', '深/实/启', '时间（估算）', '跳转链接'],
          rows, [6, 11, 7, 10, 18, 30, 60, 70, 70, 40, 6, 9, 18, 30], link_cols=(13,))

    rows = []
    for th in THEMES:
        for x in sorted((x for x in items if th in x['themes']), key=lambda x: (x['year'], x['id'])):
            rows.append([th, x['year'], x['id'], x['title'], x['speakers'], x['score'], x['link']])
    sheet(wb.create_sheet('主题索引'), ['主题', '年份', '编号', '标题', '发言人', '总分', '跳转链接'],
          rows, [14, 7, 11, 50, 10, 6, 30], link_cols=(6,))

    rows = []
    for key, name in CLUSTERS.items():
        group = sorted((x for x in items if x['cluster'] == key), key=lambda x: (x['year'], x['id']))
        if not group:
            continue
        rows.append([name, f'{len(group)} 条，{group[0]["year"]}–{group[-1]["year"]}', notes.get(key, ''), '', '', ''])
        for x in group:
            rows.append([name, x['year'], x['title'], x['lesson'], x['id'], x['link']])
    sheet(wb.create_sheet('思想演变'), ['概念簇', '年份', '标题 / 演变综述', '投资启示', '编号', '跳转链接'],
          rows, [22, 14, 60, 50, 11, 30], link_cols=(5,))
    wb.save(path)


def load_notes(path):
    notes = {}
    if path.exists():
        for ln in path.read_text(encoding='utf-8').split('\n'):
            if '\t' in ln and not ln.startswith('#'):
                k, v = ln.split('\t', 1)
                notes[k.strip()] = v.strip()
    return notes


def main():
    chapters = load_chapters(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / 'output'
    entries = parse_entries(ROOT / 'entries')
    errs = validate(entries, chapters)
    if errs:
        print('\n'.join(errs))
        sys.exit(f'{len(errs)} problems; fix entries first')
    urls = load_urls(ROOT / 'data' / 'cnbc_urls.tsv')
    notes = load_notes(ROOT / 'data' / 'cluster_notes.tsv')
    items = enrich(entries, chapters, urls)
    top = pick_top(items)
    catalog = catalog_rows(urls, chapters)
    out.mkdir(parents=True, exist_ok=True)
    write_markdown(out, items, top, catalog, notes)
    write_excel(out / '巴芒股东大会精华数据库.xlsx', items, top, catalog, notes)
    slim = [{k: v for k, v in x.items() if k not in ('quotes', 'rel_chapters', 'file')} for x in items]
    (out / 'entries.json').write_text(json.dumps({
        'entries': slim, 'top': [x['id'] for x in top], 'catalog': catalog,
        'themes': THEMES, 'clusters': CLUSTERS, 'notes': notes,
    }, ensure_ascii=False), encoding='utf-8')
    print(f'{len(items)} entries, top {len(top)}; '
          f"years {Counter(x['year'] for x in top).most_common(3)}; "
          f"clusters {Counter(x['cluster'] for x in top).most_common(3)}")


if __name__ == '__main__':
    main()
